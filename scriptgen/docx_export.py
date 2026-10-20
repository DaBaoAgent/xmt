"""把逐字稿导出成可直接打印的 Word。

只包含：标题 + 账号/出镜人 + 口播逐字稿（或双人对白稿）。
不含：对标拆解、拍摄提示、合规说明、发布包、版本记录。
"""
import os

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


FONT = "微软雅黑"


def _set_run_font(run, size, bold=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _lines(script):
    body = script.get("body") or {}
    out = []
    for line in body.get("lines") or []:
        if not isinstance(line, dict):
            continue
        text = str(line.get("text") or "").strip()
        if not text:
            continue
        out.append((str(line.get("speaker") or "").strip(), text))
    if not out:
        text = body.get("plain_text") or body.get("text") or str(body or "")
        out = [("", p.strip()) for p in str(text).splitlines() if p.strip()]
    return out


def export(script, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(16)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)

    title = str(script.get("title") or script.get("benchmark_title") or "口播逐字稿").strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(10)
    _set_run_font(p.add_run(title), 10.5, bold=True)

    account = str(script.get("account_label") or "").strip()
    presenters = [str(x) for x in (script.get("presenters") or []) if x]
    meta = " ｜ ".join(x for x in [account, "、".join(presenters)] if x)
    if meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(22)
        _set_run_font(p.add_run(meta), 10.5, color=(0x66, 0x66, 0x66))

    body = script.get("body") or {}
    is_dialogue = (body.get("format") == "dialogue") or (str(script.get("account_id")) == "brand")
    for speaker, text in _lines(script):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(8)
        if is_dialogue and speaker:
            p.paragraph_format.left_indent = Cm(0)
            _set_run_font(p.add_run(speaker + "："), 10.5, bold=True)
            _set_run_font(p.add_run(text), 10.5)
        else:
            p.paragraph_format.first_line_indent = Pt(32)
            _set_run_font(p.add_run(text), 10.5)

    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path
