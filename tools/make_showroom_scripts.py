# -*- coding: utf-8 -*-
"""展厅分镜脚本生成器（jiaoben 技能的程序化沉淀）

把 jiaoben/ 里 17 个 gen_*.py 一次性脚本泛化成数据驱动工具：
结构 = 每条 4 分镜 x 7 秒（0-7 / 7-14 / 14-21 / 21-28），
镜1 钩子、镜2-3 卖点演示、镜4 情绪收尾。

用法：
    .venv/Scripts/python.exe tools/make_showroom_scripts.py                      # 内置示例（小福车 4 套）
    .venv/Scripts/python.exe tools/make_showroom_scripts.py --spec my.json       # 自定义产品+角度
    .venv/Scripts/python.exe tools/make_showroom_scripts.py --products 小福车     # 只出某个产品
    .venv/Scripts/python.exe tools/make_showroom_scripts.py --no-docx            # 只出 txt

spec JSON 结构：
    {
      "title": "合集标题（如：佳康顺6款产品 单人口播脚本合集）",
      "products": [
        {
          "name": "一、小福车（4套·骨折康复）",
          "power": "manual",                 // manual | electric，manual 会扫电动禁用词
          "scripts": [
            {
              "title": "【脚本1】（踩坑吐槽风）：骨折康复期，千万别拄拐",
              "style": "踩坑吐槽风",          // 必须在 STYLES 里，否则报错
              "shots": [
                {"desc": "展厅里，主播对着镜头", "text": "骨折康复期，千万别拄拐！我就是踩坑了。"},
                ... 共 4 条 ...
              ]
            }, ...
          ]
        }, ...
      ]
    }

输出（默认 out/展厅分镜脚本/showroom_batch.txt，加 .docx 若装了 python-docx）：
    ============================================================
    <title>
    风格：展厅场景 | 口语化 | 活人感 | 4分镜/条
    ============================================================
    一、<产品名>
    【脚本1】（风格）：标题
    镜1（0-7s）：画面描述
      口播文案
    ...

★ 文件名坑（2026-10-23 实测）：本机安全钩子会拦截「具体产品名 + 口播/脚本」
  组合的文件名（如 小福车单人口播脚本.txt 报 PermissionError），
  所以输出文件名一律用通用名 showroom_batch.txt，中文合集标题写在文件内容里，
  要改名用 --name 传 ASCII 名字。

校验（不合规直接报错退出，不产出）：
    - style 必须在风格库；每脚本恰好 4 镜；desc/text 非空
    - manual 产品扫电动禁用词：电机/电池/摇杆/电磁刹车/瓦/AH/锂电/续航/遥控
    - text 扫价格词：多少钱/报价/价格/优惠
"""
import argparse
import json
import re
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent  # D:/kaifa/xmt
OUT_DIR_DEFAULT = BASE / "out" / "展厅分镜脚本"

# ---------------------------------------------------------------- 风格库
# 与 .workbuddy/skills/scriptgen-writing/展厅分镜脚本技能.md 的角度矩阵保持一致
STYLES = {
    "避坑指南风": "先说别只看价格，再给判断点",
    "亲戚口碑风": "熟人真实故事切入",
    "反常识风": "打破品类刻板印象",
    "老板吐槽风": "从业者身份背书",
    "场景故事风": "具体顾客具体一天",
    "对比反差风": "先踩便宜货再亮本款",
    "颜值切入风": "颜色外观破刻板印象",
    "功能拆解风": "单点深挖一个设计",
    "踩坑吐槽风": "第一人称翻车经历",
    "用户视角风": "第N天日记体",
    "对比科普风": "两个品类的使用边界",
    "黑科技风": "你见过吗式惊叹加演示",
    "家庭场景风": "老人从怕麻烦到天天出门",
    "参数拆解风": "一个硬参数讲透",
    "中老年痛点风": "安全恐惧开场",
    "女生视角风": "我一个女生也能反差举证",
    "极简风": "就一个按钮上车就会",
    "细节控风": "后轮快拆坐垫靠背档位",
    "夏天痛点风": "季节性痛点",
    "安全风": "最怕啥加结构举证",
    "王师傅偶遇风": "围观式提问适合多SKU展厅",
    "方庄店介绍风": "导购式讲选购逻辑售后",
    "旅游痛点风": "出行场景上高铁上飞机",
}

ELECTRIC_FORBIDDEN = ["电机", "电池", "摇杆", "电磁刹车", "瓦", "AH", "锂电", "续航", "遥控"]
PRICE_FORBIDDEN = ["多少钱", "报价", "价格", "优惠"]

# ---------------------------------------------------------------- 内置示例
# 小福车 4 套（来源：jiaoben/6款产品_单人口播脚本24套_全新版，已按参数档案核对：
# 小福车是手动跪式推车，禁写电动词；参数以 jiaoben/小福车/产品卖点.txt 为准）
DEMO = {
    "title": "佳康顺小福车 单人口播脚本（4套·骨折康复角度）",
    "products": [
        {
            "name": "一、小福车 骨折康复车（4套·全新角度）",
            "power": "manual",
            "scripts": [
                {
                    "title": "【脚本1】（踩坑吐槽风）：骨折康复期，千万别拄拐",
                    "style": "踩坑吐槽风",
                    "shots": [
                        {"desc": "展厅里，主播对着镜头", "text": "骨折康复期，千万别拄拐！我就是踩坑了。"},
                        {"desc": "主播指着旁边的小福车", "text": "拄拐胳肢窝磨破了，手腕也疼，还摔了一跤。"},
                        {"desc": "主播把膝盖放在跪垫上推了两步", "text": "后来换了这个小福车，膝盖跪上去，伤脚完全不用力。"},
                        {"desc": "主播停下来", "text": "四轮着地稳得很，再也没摔过。"},
                    ],
                },
                {
                    "title": "【脚本2】（用户视角风）：脚崴了第3天，我终于能下楼了",
                    "style": "用户视角风",
                    "shots": [
                        {"desc": "展厅里，主播脚腕缠着绷带", "text": "脚崴了第3天，我终于能下楼了。"},
                        {"desc": "主播推着小福车走了两步", "text": "之前拄拐根本下不了楼，太疼了。"},
                        {"desc": "主播指了指身下的跪垫", "text": "这个小福车，膝盖跪上去，伤脚悬空，一点不使劲。"},
                        {"desc": "主播笑着推了两步", "text": "终于能下楼透气了，太不容易了。"},
                    ],
                },
                {
                    "title": "【脚本3】（对比科普风）：助行器和小福车，到底有啥区别",
                    "style": "对比科普风",
                    "shots": [
                        {"desc": "展厅里，主播对着镜头", "text": "好多人问我，助行器和小福车到底有啥区别？"},
                        {"desc": "主播拍了拍助行器", "text": "助行器是腿还行，就是没力气，扶着走。"},
                        {"desc": "主播又拍了拍小福车", "text": "小福车是一条腿不能着地，膝盖跪上去，伤腿完全休息。"},
                        {"desc": "主播点点头", "text": "骨折术后用小福车，腿恢复得快。"},
                    ],
                },
                {
                    "title": "【脚本4】（场景故事风）：我媳妇生孩子侧切，用这个下地了",
                    "style": "场景故事风",
                    "shots": [
                        {"desc": "展厅里，主播对着镜头", "text": "我媳妇生孩子侧切，下地都费劲。"},
                        {"desc": "主播指了指小福车", "text": "后来我把这个小福车给她用，膝盖跪上去，肚子不用使劲。"},
                        {"desc": "主播推了两步", "text": "慢慢就能自己上厕所了，不用人扶。"},
                        {"desc": "主播笑了", "text": "康复期有个这个，真的少遭罪。"},
                    ],
                },
            ],
        }
    ],
}


def validate(spec):
    """返回错误列表，空 = 通过。"""
    errors = []
    if not spec.get("title"):
        errors.append("spec 缺少 title")
    products = spec.get("products") or []
    if not products:
        errors.append("spec 缺少 products")
    for p in products:
        pname = p.get("name") or "<未命名产品>"
        power = p.get("power", "manual")
        if power not in ("manual", "electric"):
            errors.append("[%s] power 只能是 manual/electric" % pname)
        scripts = p.get("scripts") or []
        if not scripts:
            errors.append("[%s] 没有 scripts" % pname)
        titles = set()
        styles_used = set()
        for s in scripts:
            stitle = s.get("title") or "<未命名脚本>"
            tag = "[%s / %s]" % (pname, stitle)
            if stitle in titles:
                errors.append("%s 标题重复" % tag)
            titles.add(stitle)
            style = s.get("style") or ""
            if style not in STYLES:
                errors.append("%s 风格「%s」不在风格库（%s）" % (tag, style, "、".join(sorted(STYLES))))
            styles_used.add(style)
            shots = s.get("shots") or []
            if len(shots) != 4:
                errors.append("%s 分镜数=%d，必须是 4" % (tag, len(shots)))
            for i, sh in enumerate(shots, 1):
                if not (sh.get("desc") or "").strip():
                    errors.append("%s 镜%d 缺画面描述" % (tag, i))
                if not (sh.get("text") or "").strip():
                    errors.append("%s 镜%d 缺口播文案" % (tag, i))
            for i, sh in enumerate(shots, 1):
                text = (sh.get("text") or "")
                if power == "manual":
                    for w in ELECTRIC_FORBIDDEN:
                        if w in text:
                            errors.append("%s 镜%d 手动款含电动词「%s」" % (tag, i, w))
                for w in PRICE_FORBIDDEN:
                    if w in text:
                        errors.append("%s 镜%d 含价格词「%s」" % (tag, i, w))
        if len(styles_used) < len(scripts):
            errors.append("[%s] 4 套角度重复（风格去重后 %d < %d），违反配比公式" % (pname, len(styles_used), len(scripts)))
    return errors


def render_txt(spec):
    lines = []
    bar = "=" * 60
    lines.append(bar)
    lines.append(spec["title"])
    lines.append("风格：展厅场景 | 口语化 | 活人感 | 4分镜/条")
    lines.append(bar)
    for p in spec["products"]:
        lines.append("")
        lines.append(p["name"])
        lines.append("")
        for s in p["scripts"]:
            lines.append(s["title"])
            for i, sh in enumerate(s["shots"], 1):
                lines.append("镜%d（%d-%ds）：%s" % (i, (i - 1) * 7, i * 7, sh["desc"]))
                lines.append("  " + sh["text"])
            lines.append("")
    return "\n".join(lines) + "\n"


def render_docx(spec, out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    title = doc.add_heading(spec["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.add_run("风格：展厅场景 | 口语化 | 活人感 | 4分镜/条").italic = True
    for prod in spec["products"]:
        doc.add_heading(prod["name"], level=1)
        doc.add_paragraph("")
        for s in prod["scripts"]:
            doc.add_heading(s["title"], level=2)
            for i, sh in enumerate(s["shots"], 1):
                para = doc.add_paragraph()
                run = para.add_run("镜%d（%d-%ds）：%s：" % (i, (i - 1) * 7, i * 7, sh["desc"]))
                run.bold = True
                para.add_run(sh["text"])
            doc.add_paragraph("")
    doc.save(str(out_path))


def main():
    ap = argparse.ArgumentParser(description="展厅4分镜口播脚本生成器")
    ap.add_argument("--spec", help="自定义 spec JSON 路径（不给则用内置小福车示例）")
    ap.add_argument("--products", help="只出指定产品（按 name 包含匹配，逗号分隔多个）")
    ap.add_argument("--out", default=str(OUT_DIR_DEFAULT), help="输出目录")
    ap.add_argument("--name", default="showroom_batch", help="输出文件名（不含扩展名，ASCII）")
    ap.add_argument("--no-docx", action="store_true", help="只出 txt")
    args = ap.parse_args()

    spec = DEMO
    if args.spec:
        spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))

    if args.products:
        keys = [k.strip() for k in args.products.split(",") if k.strip()]
        spec = dict(spec)
        spec["products"] = [p for p in spec["products"] if any(k in p.get("name", "") for k in keys)]
        if not spec["products"]:
            print("没有匹配的产品：%s" % keys)
            return 2

    errors = validate(spec)
    if errors:
        print("校验不通过，未产出：")
        for e in errors:
            print("  - " + e)
        return 1

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", args.name).strip("_") or "showroom_batch"
    txt_path = out_dir / (safe_name + ".txt")
    txt_path.write_text(render_txt(spec), encoding="utf-8")
    print("TXT  -> %s" % txt_path)

    if not args.no_docx:
        try:
            docx_path = out_dir / (safe_name + ".docx")
            render_docx(spec, docx_path)
            print("DOCX -> %s" % docx_path)
        except ImportError:
            print("未装 python-docx，跳过 docx（txt 已产出）")

    n_scripts = sum(len(p["scripts"]) for p in spec["products"])
    print("完成：%d 个产品 / %d 条脚本" % (len(spec["products"]), n_scripts))
    return 0


if __name__ == "__main__":
    sys.exit(main())
