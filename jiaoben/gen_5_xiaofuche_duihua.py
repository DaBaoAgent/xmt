# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def add_title(text, level=1):
    p = doc.add_heading(text, level=level)
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = '微软雅黑'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return p

def add_para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(10.5)
    run.bold = bold
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = '微软雅黑'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                run.font.size = Pt(9)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = '微软雅黑'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                    run.font.size = Pt(9)
    return table

add_title('小福车 双人对话脚本5套', level=0)
add_title('（皮皮熊风格·公园场景·一男一女）', level=0)
doc.add_paragraph()
add_para('产品：小福车（膝盖跪垫式康复车）', bold=True)
add_para('人群：脚踝骨折/术后康复期 | 演员：小宇（男）、小雅（女）', bold=True)
doc.add_page_break()

scripts = [
    {
        "name": "脚本1",
        "title": "脚崴了拄拐太疼了，试试这个",
        "cover_main": "拄拐太疼了",
        "cover_sub": "试试这个",
        "rows": [
            ["1", "0-7s", "公园步道，小宇脚腕缠着绷带，拄着拐杖走，一脸痛苦。小雅推着小福车过来", "小宇：哎哟，拄拐真的太疼了，胳肢窝都磨破了。小雅：别拄拐了，试试这个。"],
            ["2", "7-14s", "小雅把小福车停在小宇面前，指了指跪垫", "小雅：这是小福车，膝盖跪上去，伤脚完全不用着地。小宇：膝盖跪着？能舒服吗？"],
            ["3", "14-21s", "小宇半信半疑跪上去，推了两步，眼睛亮了", "小宇：哎？真的不用使劲，四轮稳得很。小雅：对吧，还有减震弹簧，不颠。"],
            ["4", "21-28s", "小宇推着车逛了两步，笑了", "小宇：这比拄拐舒服一百倍！小雅：康复期用这个，少遭多少罪。"],
        ]
    },
    {
        "name": "脚本2",
        "title": "你脚骨折了怎么还能逛公园？",
        "cover_main": "脚骨折了",
        "cover_sub": "怎么还逛公园",
        "rows": [
            ["1", "0-7s", "公园步道，小雅推着小福车，膝盖跪在上面，小宇跟在旁边走", "小宇：你脚不是骨折了吗？怎么还能逛公园？"],
            ["2", "7-14s", "小雅推了两步，指了指身下的小福车", "小雅：你看这个，小福车，膝盖跪垫式的，伤脚完全悬空不用力。"],
            ["3", "14-21s", "小雅在小坡上停住，捏了捏刹车", "小雅：四轮着地稳得很，手刹一捏就停，逛公园完全没问题。小宇：这也太方便了吧。"],
            ["4", "21-28s", "小雅笑着推走了", "小雅：康复期也能出门，不用在家闷着。小宇：给我也整一个。"],
        ]
    },
    {
        "name": "脚本3",
        "title": "拄拐摔了两次，换这个就没摔过",
        "cover_main": "拄拐摔了两次",
        "cover_sub": "换这个就没摔过",
        "rows": [
            ["1", "0-7s", "公园长椅旁，小雅坐着揉脚，小宇在旁边扶着小福车", "小宇：你脚好点没？之前拄拐不是摔了两次吗？"],
            ["2", "7-14s", "小雅拍了拍旁边的小福车", "小雅：别提了，拄拐又疼又不稳，摔了两次。后来换了这个小福车。"],
            ["3", "14-21s", "小雅跪上去推了两步，很稳", "小雅：四轮着地，怎么推都不晃，还能折叠放后备箱。小宇：那是稳多了。"],
            ["4", "21-28s", "小雅停下来笑了", "小雅：复查遛弯都靠它，康复期舒服多了。小宇：早用这个就好了。"],
        ]
    },
    {
        "name": "脚本4",
        "title": "脚踝骨折三个月，全靠它撑过来",
        "cover_main": "骨折三个月",
        "cover_sub": "全靠它撑过来",
        "rows": [
            ["1", "0-7s", "公园步道，小宇和小雅一起走，小雅推着小福车", "小宇：你脚踝骨折三个月了，恢复得还不错啊？"],
            ["2", "7-14s", "小雅拍了拍小福车", "小雅：全靠它啊，一开始拄拐疼得受不了，后来用了这个小福车。"],
            ["3", "14-21s", "小雅展示折叠，抱起来走了两步", "小雅：膝盖跪着走，伤脚不用力，还能折叠，放后备箱去复查特方便。小宇：这也太实用了。"],
            ["4", "21-28s", "两人笑着往前走", "小雅：康复期选对工具，真的少走好多弯路。小宇：没错，舒服多了。"],
        ]
    },
    {
        "name": "脚本5",
        "title": "别再拄拐了，试试这个",
        "cover_main": "别再拄拐了",
        "cover_sub": "试试这个",
        "rows": [
            ["1", "0-7s", "公园空地，小宇拄着拐杖站着，一脸痛苦。小雅推着小福车过来", "小宇：哎，拄拐真的太累了，胳肢窝疼，走路还晃。"],
            ["2", "7-14s", "小雅把小福车推到小宇面前", "小雅：别拄拐了，试试这个小福车。小宇：这是什么？"],
            ["3", "14-21s", "小雅示范跪在上面推了两步", "小雅：膝盖跪垫式的，伤脚完全悬空，四轮稳得很，还有减震弹簧。"],
            ["4", "21-28s", "小宇跪上去试了试，眼睛亮了", "小宇：哎？真的舒服多了！小雅：对吧，康复期用这个，少遭罪。"],
        ]
    },
]

for i, s in enumerate(scripts):
    add_title(s["name"] + "：" + s["title"], level=1)
    add_para('封面：' + s["cover_main"] + " / " + s["cover_sub"], bold=True)
    doc.add_paragraph()
    add_table(['镜号', '时长', '画面', '台词'], s["rows"])
    if i < len(scripts) - 1:
        doc.add_page_break()

output_path = r'D:\kaifa\jiaoben\小福车双人对话脚本5套.docx'
doc.save(output_path)
print('文档已生成：' + output_path)
