# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def add_title(text, level=1):
    p = doc.add_heading(text, level=level)
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = '微软雅黑'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return p

def add_para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(10.5)
    run.bold = bold
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = '微软雅黑'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                run.font.size = Pt(9)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = '微软雅黑'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                    run.font.size = Pt(9)
    return table

add_title('小福车 皮皮熊风格脚本5套', level=0)
add_title('（情感痛点+场景展示·公园场景）', level=0)
doc.add_paragraph()
add_para('产品：小福车（膝盖跪垫式康复车）', bold=True)
add_para('人群：脚踝骨折/跖骨骨折/跟腱术后/脚扭伤康复期', bold=True)
doc.add_page_break()

scripts = [
    {
        "name": "脚本1",
        "title": "脚崴了不用拄拐，这个太舒服了",
        "cover_main": "脚崴了不用拄拐",
        "cover_sub": "这个太舒服了",
        "rows": [
            ["1", "0-7s", "公园步道，小雅扶着小福车慢慢走，脚腕还缠着绷带", "脚崴了拄拐真的太难受了，胳肢窝都磨破了。"],
            ["2", "7-14s", "小雅把膝盖放在跪垫上，推着走了两步", "后来发现这个小福车，膝盖跪上去，伤脚完全不用着地。"],
            ["3", "14-21s", "小雅推过一段不平的路，很稳", "四轮着地稳得很，还有减震弹簧，推着不颠。"],
            ["4", "21-28s", "小雅停下来，笑了一下", "康复期用这个，比拄拐舒服太多了。"],
        ]
    },
    {
        "name": "脚本2",
        "title": "以前拄拐磨得疼，现在用这个逛公园",
        "cover_main": "以前拄拐磨得疼",
        "cover_sub": "现在用这个逛公园",
        "rows": [
            ["1", "0-7s", "公园长椅旁，小宇对着镜头，小雅在旁边扶着小福车", "我脚骨折那阵，拄拐拄得胳肢窝都破了，出门都不想出。"],
            ["2", "7-14s", "小雅示范跪在小福车上，推了两步", "后来用了这个小福车，膝盖跪着走，伤脚完全悬空。"],
            ["3", "14-21s", "小雅推到小宇旁边，小宇坐上去试了试", "你看，四轮稳得很，转弯也灵活，逛公园完全没问题。"],
            ["4", "21-28s", "两人推着车笑了", "康复期有这个，真的太省心了。"],
        ]
    },
    {
        "name": "脚本3",
        "title": "注意看，这个女生脚扭了居然还能逛公园",
        "cover_main": "注意看这个女生",
        "cover_sub": "脚扭了还逛公园",
        "rows": [
            ["1", "0-7s", "公园步道，小雅推着小福车，膝盖跪在上面，对着镜头", "注意看，这个女生脚扭了，居然还能逛公园。"],
            ["2", "7-14s", "小雅推了两步，转了个弯", "秘密就是这个小福车，膝盖跪垫式的，伤脚完全不用着地。"],
            ["3", "14-21s", "小雅指了指前面的袋子", "前面还能放手机钥匙，手刹一捏就停，特安全。"],
            ["4", "21-28s", "小雅笑着推走了", "康复期也能出门，真的太方便了。"],
        ]
    },
    {
        "name": "脚本4",
        "title": "脚踝骨折三个月，全靠它撑过来",
        "cover_main": "脚踝骨折三个月",
        "cover_sub": "全靠它撑过来",
        "rows": [
            ["1", "0-7s", "公园步道，小雅推着小福车对着镜头", "脚踝骨折那三个月，我全靠这个小福车撑过来。"],
            ["2", "7-14s", "小雅跪在上面推了两步", "一开始拄拐，磨得胳肢窝疼，还摔过两次。"],
            ["3", "14-21s", "小雅展示折叠，抱起来走了两步", "后来换了这个，膝盖跪着走，四轮稳得很，还能折叠放后备箱。"],
            ["4", "21-28s", "小雅笑了", "复查遛弯都靠它，真的是康复神器。"],
        ]
    },
    {
        "name": "脚本5",
        "title": "选对康复工具，少走半年弯路",
        "cover_main": "选对康复工具",
        "cover_sub": "少走半年弯路",
        "rows": [
            ["1", "0-7s", "公园长椅旁，小雅对着镜头", "脚骨折了选对康复工具，真的少走半年弯路。"],
            ["2", "7-14s", "小雅拍了拍旁边的小福车", "一开始我也以为只能拄拐，后来才知道有这个小福车。"],
            ["3", "14-21s", "小雅跪在上面推了两步", "膝盖跪垫式，伤脚完全悬空，四轮稳，还能折叠。"],
            ["4", "21-28s", "小雅点点头", "康复期用这个，舒服又安全。"],
        ]
    },
]

for i, s in enumerate(scripts):
    add_title(s["name"] + "：" + s["title"], level=1)
    add_para('封面：' + s["cover_main"] + " / " + s["cover_sub"], bold=True)
    doc.add_paragraph()
    add_table(['镜号', '时长', '画面', '台词'], s["rows"])
    if i < len(scripts) - 1:
        doc.add_page_break()

output_path = r'D:\kaifa\jiaoben\小福车皮皮熊风格脚本5套.docx'
doc.save(output_path)
print('文档已生成：' + output_path)
