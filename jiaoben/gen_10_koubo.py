# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def add_title(text, level=1):
    p = doc.add_heading(text, level=level)
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = '微软雅黑'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return p

def add_para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(10.5)
    run.bold = bold
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = '微软雅黑'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                run.font.size = Pt(9)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = '微软雅黑'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                    run.font.size = Pt(9)
    return table

add_title('朋友2助行器 单人嘴播脚本10套', level=0)
add_title('（公园场景·口语化·4分镜版）', level=0)
doc.add_paragraph()
add_para('形式：一个人对着镜头口播，边演示边说', bold=True)
add_para('场景：公园 | 时长：25-30秒/条 | 风格：像跟朋友聊天，不生硬', bold=True)
doc.add_page_break()

scripts = [
    {
        "name": "脚本1：猜产品",
        "title": "你们猜这是啥？真不是购物车",
        "cover_main": "你们猜这是啥",
        "cover_sub": "真不是购物车",
        "rows": [
            ["1", "0-7s", "公园步道，小雅推着车对着镜头走过来，拍了拍车把", "哎你们猜这是啥？是不是以为购物车？我一开始也以为是。"],
            ["2", "7-14s", "镜头特写，小雅一翻座椅坐上去，又翻回来", "你看啊，座椅一翻就能当轮椅推，翻过来自己扶着走就行，一个顶俩。"],
            ["3", "14-21s", "小雅往后面挂了个袋子，又单手拎起来走了两步", "后面能挂个袋子，买个菜啥的直接推着走。而且贼轻，才7.8公斤，我单手就拎起来了。"],
            ["4", "21-28s", "小雅对着镜头笑了下，推着车往前走", "最主要这绿颜色真好看，推出去一点都不显老。"],
        ]
    },
    {
        "name": "脚本2：重量轻",
        "title": "才7.8公斤，我单手就拎起来了",
        "cover_main": "才7.8公斤",
        "cover_sub": "单手就拎起来",
        "rows": [
            ["1", "0-7s", "公园空地，小雅对着镜头，旁边放着折叠好的助行器", "你们见过助行器吗？都老沉老沉的对吧？你看我这个。"],
            ["2", "7-14s", "小雅单手拎起助行器，举了两下", "才7.8公斤，我一只手就拎起来了，比你那电脑包重不了多少。"],
            ["3", "14-21s", "小雅把车放地上，一按就展开了", "展开也贼简单，一按就开，老人自己都能弄。"],
            ["4", "21-28s", "小雅推着走了两步", "轻才是真的好，老人提着不累。"],
        ]
    },
    {
        "name": "脚本3：折叠方便",
        "title": "一拉就折好，塞后备箱正正好",
        "cover_main": "一拉就折好",
        "cover_sub": "塞后备箱正正好",
        "rows": [
            ["1", "0-7s", "公园长椅旁，小雅对着镜头", "带老人出门最烦啥？助行器太大塞不进车对吧？你看我这个。"],
            ["2", "7-14s", "小雅掀开坐垫，一拉就折好了", "掀个坐垫一拉就折，就这么简单，女生都能弄。"],
            ["3", "14-21s", "小雅把折叠好的车比了比大小", "折完就这么点，往公园角落一放也不占地方，放后备箱正正好。"],
            ["4", "21-28s", "小雅拍了拍车", "出门带着一点不费劲。"],
        ]
    },
    {
        "name": "脚本4：反转座椅",
        "title": "翻一下就能当轮椅，这个设计真绝了",
        "cover_main": "翻一下就能当轮椅",
        "cover_sub": "设计真绝了",
        "rows": [
            ["1", "0-7s", "公园步道，小雅推着助行器对着镜头", "你们有没有想过，助行器还能当轮椅推？"],
            ["2", "7-14s", "小雅把座椅翻过来，自己坐上去", "你看啊，座椅这么一翻，我坐上去，你就能推着我走。"],
            ["3", "14-21s", "小雅又翻回去，扶着扶手站着", "翻回来呢，我自己扶着走，锻炼也行。一车两用，不用单独买轮椅了。"],
            ["4", "21-28s", "小雅推了两步", "这个反转设计，真的贼聪明。"],
        ]
    },
    {
        "name": "脚本5：颜色好看",
        "title": "推出去谁能想到这是助行器啊",
        "cover_main": "谁能想到这是助行器",
        "cover_sub": "颜值太高了",
        "rows": [
            ["1", "0-7s", "公园步道，小雅推着珍珠绿的车对着镜头转了一圈", "你们看这个颜色，是不是特别好看？谁能想到这是个助行器啊。"],
            ["2", "7-14s", "镜头给车身漆面一个特写", "就这个绿，还有个紫色的，特别显年轻，一点都不像那种老气的银色。"],
            ["3", "14-21s", "小雅摸了摸车架", "做工也细，不是那种便宜货，推出去真有面儿。"],
            ["4", "21-28s", "小雅笑着推走了", "老人用这个，心情都好了。"],
        ]
    },
    {
        "name": "脚本6：买菜神器",
        "title": "买菜别再手拎了，推着它就走",
        "cover_main": "买菜别再手拎了",
        "cover_sub": "推着它就走",
        "rows": [
            ["1", "0-7s", "公园长椅旁，小雅挂着个装满东西的袋子，对着镜头", "你们买菜是不是都得手拎着？勒得手通红对吧？"],
            ["2", "7-14s", "小雅拍了拍后面的购物袋", "你看这个，标配购物袋，买完菜直接挂后面，推着就走，手上啥也不用拿。"],
            ["3", "14-21s", "小雅又指了指车上的水杯架", "还有水杯架，水放这儿，不用手拿着，特方便。"],
            ["4", "21-28s", "小雅推着走了两步", "老人买菜用这个，太省事了。"],
        ]
    },
    {
        "name": "脚本7：稳得很",
        "title": "过个坎儿都不晃，稳得像坦克",
        "cover_main": "过坎儿都不晃",
        "cover_sub": "稳得像坦克",
        "rows": [
            ["1", "0-7s", "公园不平的小路上，小雅推着车对着镜头", "助行器最怕啥？不稳对吧？推个坎儿就晃。你看我这个。"],
            ["2", "7-14s", "小雅推着车过了一个小坎儿，车很稳", "看到没？就这么稳，金字塔结构的，分散受力，怎么推都不晃。"],
            ["3", "14-21s", "小雅推过一段碎石路", "碎石路也不颠，轮子顺滑得很，老人用着不颠屁股。"],
            ["4", "21-28s", "小雅点点头", "稳才是真的安全。"],
        ]
    },
    {
        "name": "脚本8：赠品全",
        "title": "买个助行器送一堆东西，太良心了",
        "cover_main": "送一堆东西",
        "cover_sub": "太良心了",
        "rows": [
            ["1", "0-7s", "公园长椅上摆着一堆配件：坐垫、购物袋、水杯架、拐杖支架，小雅对着镜头", "你们买助行器都送啥？你看我这个，送了一堆。"],
            ["2", "7-14s", "小雅拿起坐垫拍了拍", "加厚坐垫，坐着不硌。购物袋，买菜用的。"],
            ["3", "14-21s", "小雅又拿起水杯架和拐杖支架", "还有水杯架、拐杖支架，都给你配好了，不用再单独买。"],
            ["4", "21-28s", "小雅把东西都装上车", "这个价位配这么全，真的挺良心的。"],
        ]
    },
    {
        "name": "脚本9：刹车安全",
        "title": "一捏就停，上下坡不溜车",
        "cover_main": "一捏就停",
        "cover_sub": "上下坡不溜",
        "rows": [
            ["1", "0-7s", "公园小坡上，小雅对着镜头", "助行器最危险啥？上下坡溜车对吧？你看我这个刹车。"],
            ["2", "7-14s", "小雅推着车到坡上，一捏刹车，车稳稳停住", "看到没？一捏就停，坡上也不溜，就这么稳。"],
            ["3", "14-21s", "小雅松开刹车，车又往前走了", "松手就走，捏紧就停，想停就停，特别好控制。"],
            ["4", "21-28s", "小雅点点头", "老人用这个，安全才是最重要的。"],
        ]
    },
    {
        "name": "脚本10：别买笨重轮椅",
        "title": "别再给爸妈买那种笨重轮椅了",
        "cover_main": "别买笨重轮椅了",
        "cover_sub": "现在都用这个",
        "rows": [
            ["1", "0-7s", "公园空地，小雅对着镜头", "你们是不是还给爸妈买那种大轮椅？老沉老沉的，搬都搬不动。"],
            ["2", "7-14s", "小雅拍了拍旁边小巧的朋友2", "别买那个了，你看这个，才7.8公斤，比那轮椅轻一半都不止。"],
            ["3", "14-21s", "小雅翻下座椅，又折起来", "能推能坐，一拉就折，女生都能拎。"],
            ["4", "21-28s", "小雅笑着推走了", "老人用着才方便。"],
        ]
    },
]

for i, s in enumerate(scripts):
    add_title(s["name"] + "：" + s["title"], level=1)
    add_para('封面：' + s["cover_main"] + " / " + s["cover_sub"], bold=True)
    doc.add_paragraph()
    add_table(['镜号', '时长', '画面', '口播台词'], s["rows"])
    if i < len(scripts) - 1:
        doc.add_page_break()

output_path = r'D:\kaifa\jiaoben\朋友2助行器单人嘴播脚本10套.docx'
doc.save(output_path)
print('文档已生成：' + output_path)
